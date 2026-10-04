# word_agent.py
# WordReaderAgent - reads paragraphs, headings and tables from .docx files

import docx

from agents.base_agent import BaseAgent


class WordReaderAgent(BaseAgent):
    name = "WordReaderAgent"
    extensions = [".docx"]

    def read_file(self, file_path):
        document = docx.Document(file_path)
        sections = []
        headings = []
        word_count = 0
        paragraph_count = 0

        # group the paragraphs under their heading
        current_heading = "Introduction"
        current_text = []

        for para in document.paragraphs:
            text = para.text.strip()
            if text == "":
                continue
            paragraph_count += 1
            word_count += len(text.split())

            style_name = ""
            if para.style is not None and para.style.name:
                style_name = para.style.name.lower()

            if style_name.startswith("heading") or style_name == "title":
                # save the previous heading block
                if len(current_text) > 0:
                    sections.append({"text": current_heading + "\n" + "\n".join(current_text),
                                     "info": {"section": current_heading[:200]}})
                current_heading = text
                current_text = []
                headings.append(text)
            else:
                current_text.append(text)

        # save the last block
        if len(current_text) > 0:
            sections.append({"text": current_heading + "\n" + "\n".join(current_text),
                             "info": {"section": current_heading[:200]}})

        # tables
        table_number = 1
        for table in document.tables:
            rows = []
            for row in table.rows:
                cells = []
                for cell in row.cells:
                    cell_text = cell.text.strip()
                    # merged cells come again and again, so skip repeated ones
                    if len(cells) == 0 or cells[-1] != cell_text:
                        cells.append(cell_text)
                rows.append(" | ".join(cells))
            if len(rows) > 0:
                sections.append({"text": f"Table {table_number}:\n" + "\n".join(rows),
                                 "info": {"section": f"Table {table_number}"}})
            table_number += 1

        profile = {
            "paragraphs": paragraph_count,
            "tables": len(document.tables),
            "word_count": word_count,
            "headings": headings[:50],
            "title": document.core_properties.title or None,
            "author": document.core_properties.author or None,
        }
        return sections, profile, []
